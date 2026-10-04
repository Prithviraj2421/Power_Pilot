"""Synthetic "legacy reports" with known formulas and planted mistakes, built from each benchmark dataset.

Per dataset up to two reports (an .xlsx, then a .csv, which only carries shown text):
  * "SUM(m) by dim" with a typed "All" row and a formula Total row (xlsx) / typed Total row (csv)
  * "AVERAGE(m2) by dim", "rows by dim", and "SUM(m) / SUM(m2) by dim" when two measures exist
Three mistakes are planted in the first table, wherever the numbers are precise enough for a mistake to be meaningful:
  swapped-digits   two adjacent different digits of one cell are swapped
  one-row          one cell equals the true value minus a single row's value
  missing-category the typed "All" row equals the grand total minus one whole category
Everything is deterministic from the benchmark seed. The truth (formula of every correct cell, kind of every mistake) is
written beside each report as <name>.truth.json.

    python -m eval.reports.generate [dataset ids...]
"""

from __future__ import annotations

import io
import json
import sys
from dataclasses import dataclass, field
from typing import Optional

import numpy as np
import pandas as pd

from app.intelligence.kpi.compilers import PandasCompiler
from app.intelligence.kpi.ir import Compare, Expr, Filter, Measure, Op, Ratio
from app.reverse.describe import canonical_filters, ir_to_dict
from eval.common import REPORTS, SEED, Dataset, dataset_ids, load

PC = PandasCompiler()


@dataclass
class Cell:
    table: str
    row: str
    col: str
    value: float
    expected: Optional[Expr]  # None for a planted mistake or a derived total
    fmt: str = "#,##0.00"
    mistake: Optional[str] = None
    formula: Optional[str] = None


@dataclass
class ReportSpec:
    dataset: str
    name: str
    kind: str  # "xlsx" | "csv"
    tables: dict[str, list[list]] = field(default_factory=dict)
    cells: list[Cell] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # mistakes that could not be planted, and why

    def truth(self) -> dict:
        return {
            "dataset": self.dataset,
            "report": self.name,
            "skipped": self.skipped,
            "cells": [
                {
                    "table": c.table, "row": c.row, "col": c.col, "shown": shown(c), "mistake": c.mistake,
                    "derived": c.formula is not None, "formula": ir_to_dict(c.expected) if c.expected else None,
                }
                for c in self.cells
            ],
        }


def shown(cell: Cell) -> str:
    return f"{cell.value * 100:.1f}%" if cell.fmt == "0.0%" else f"{cell.value:,.2f}"


def pick_dimensions(ds: Dataset) -> list[str]:
    """Text columns with 3 to 8 values, fewest first: the natural "by region / by department" axes."""
    frame = ds.frame
    found = [
        (frame[c].nunique(), c) for c in frame.columns
        if (pd.api.types.is_string_dtype(frame[c]) or frame[c].dtype == object)
        and 3 <= frame[c].nunique() <= 8 and frame[c].notna().all()
        and all(len(str(v)) >= 2 for v in frame[c].unique())
    ]
    return [c for _, c in sorted(found)]


def pick_measures(ds: Dataset) -> list[str]:
    frame = ds.frame
    return [
        c for c in frame.columns
        if ds.columns.get(c) not in ("identifier", "date")
        and pd.api.types.is_numeric_dtype(frame[c]) and not pd.api.types.is_bool_dtype(frame[c])
        and frame[c].nunique() >= 20 and frame[c].notna().all() and frame[c].abs().sum() > 1000
    ]


def swap_digits(text: str) -> Optional[str]:
    """Swap the first two adjacent, different digits after the leading digit."""
    for i in range(1, len(text) - 1):
        if text[i].isdigit() and text[i + 1].isdigit() and text[i] != text[i + 1]:
            return text[:i] + text[i + 1] + text[i] + text[i + 2:]
    return None


def build(ds: Dataset, dim: str, measure: str, second: Optional[str], kind: str, name: str, rng: np.random.Generator) -> Optional[ReportSpec]:
    frame = ds.frame
    values = sorted(frame[dim].unique(), key=str)
    spec = ReportSpec(ds.id, name, kind)

    def of(op: Op, column: Optional[str], value: str) -> Expr:
        return Measure(op, column, (Filter(dim, Compare.EQ, value),))

    def true(expr: Expr) -> float:
        return float(PC.evaluate(expr, frame))

    # -- table 1: SUM(measure) by dim, with the three planted mistakes -------------------------------
    title = f"{measure} by {dim}"
    rows: list[list] = [[title], [dim, measure]]
    cells = []
    for value in values:
        expr = of(Op.SUM, measure, str(value))
        cells.append(Cell(title, str(value), measure, round(true(expr), 2), expr))
    order = list(rng.permutation(len(cells)))

    swapped = swap_digits(f"{cells[order[0]].value:.2f}")
    if swapped:
        cells[order[0]].value, cells[order[0]].expected, cells[order[0]].mistake = float(swapped), None, "swapped-digits"
    else:
        spec.skipped.append("swapped-digits: no two adjacent different digits to swap")

    victim = cells[order[1]]
    group = frame[frame[dim].astype(str) == victim.row][measure]
    smallest = float(group.iloc[group.abs().argmin()])
    if group.abs().min() >= 1 and abs(victim.value) > 10 * abs(smallest):
        victim.value, victim.expected, victim.mistake = round(victim.value - smallest, 2), None, "one-row"
    else:
        spec.skipped.append("one-row: no single row small enough to be a plausible omission")

    for cell in cells:
        spec.cells.append(cell)
        rows.append([cell.row, cell])
    grand = Measure(Op.SUM, measure)
    all_row = Cell(title, f"All {dim}", measure, round(true(grand), 2), grand)
    dropped = cells[order[2]] if len(cells) > 2 else None
    if dropped is not None:
        category_total = true(of(Op.SUM, measure, dropped.row))
        if abs(category_total) > 0.01 * abs(all_row.value):
            all_row.value, all_row.expected, all_row.mistake = round(all_row.value - category_total, 2), None, "missing-category"
        else:
            spec.skipped.append("missing-category: the category is too small for its absence to show")
    spec.cells.append(all_row)
    rows.append([all_row.row, all_row])
    first, last = 3, 2 + len(cells)
    total = Cell(title, "Total", measure, round(sum(c.value for c in cells), 2), None, formula=f"=SUM(B{first}:B{last})")
    spec.cells.append(total)
    rows.append(["Total", total])
    spec.tables[title] = rows

    # -- tables 2..4: no mistakes planted ----------------------------------------------------------
    others = [(pick_measures(ds)[1:2] or [None])[0] if second is None else second]
    if others[0]:
        t = f"Average {others[0]} by {dim}"
        r2: list[list] = [[t], [dim, f"Average {others[0]}"]]
        for value in values:
            expr = of(Op.AVERAGE, others[0], str(value))
            c = Cell(t, str(value), f"Average {others[0]}", round(true(expr), 2), expr)
            spec.cells.append(c)
            r2.append([c.row, c])
        spec.tables[t] = r2
        t = f"{measure} per {others[0]} by {dim}"
        r4: list[list] = [[t], [dim, f"{measure} per {others[0]}"]]
        for value in values:
            expr = Ratio(of(Op.SUM, measure, str(value)), of(Op.SUM, others[0], str(value)))
            if true(expr) == 0:
                continue
            c = Cell(t, str(value), f"{measure} per {others[0]}", round(true(expr), 2), expr)
            spec.cells.append(c)
            r4.append([c.row, c])
        spec.tables[t] = r4
    t = f"Rows by {dim}"
    r3: list[list] = [[t], [dim, "Rows"]]
    for value in values:
        expr = of(Op.COUNT, None, str(value))
        c = Cell(t, str(value), "Rows", float(true(expr)), expr, fmt="0")
        spec.cells.append(c)
        r3.append([c.row, c])
    spec.tables[t] = r3
    return spec


def to_xlsx(spec: ReportSpec) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    r = 1
    for table in spec.tables.values():
        for row in table:
            for c, item in enumerate(row, start=1):
                if isinstance(item, Cell):
                    ws.cell(r, c, item.formula or item.value).number_format = item.fmt if item.fmt != "0.0%" else "0.0%"
                else:
                    ws.cell(r, c, item)
            r += 1
        r += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_csv(spec: ReportSpec) -> bytes:
    lines = []
    for table in spec.tables.values():
        for row in table:
            cells = [shown(x) if isinstance(x, Cell) else str(x) for x in row]
            lines.append(",".join(f'"{c}"' if ("," in c or '"' in c) else c for c in cells))
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def reports_for(ds: Dataset) -> list[tuple[ReportSpec, bytes, str]]:
    """(spec, file bytes, file name) for up to two reports of one dataset; empty when the data has no usable axis."""
    dims, measures = pick_dimensions(ds), pick_measures(ds)
    if not dims or not measures:
        return []
    rng = np.random.default_rng([SEED, sum(map(ord, ds.id))])
    out = []
    plans = [(dims[0], measures[0], measures[1] if len(measures) > 1 else None, "xlsx")]
    plans.append((dims[1] if len(dims) > 1 else dims[0], measures[1] if len(measures) > 1 else measures[0], measures[0] if len(measures) > 1 else None, "csv"))
    for index, (dim, measure, second, kind) in enumerate(plans, start=1):
        spec = build(ds, dim, measure, second, kind, f"{ds.id}_report{index}", rng)
        if spec is not None:
            out.append((spec, to_xlsx(spec) if kind == "xlsx" else to_csv(spec), f"{spec.name}.{kind}"))
    return out


def main(argv: list[str]) -> int:
    REPORTS.mkdir(exist_ok=True)
    for ident in argv or dataset_ids():
        ds = load(ident)
        made = reports_for(ds)
        for spec, content, filename in made:
            (REPORTS / filename).write_bytes(content)
            (REPORTS / f"{spec.name}.truth.json").write_text(json.dumps(spec.truth(), indent=2), encoding="utf-8")
        print(f"{ident:28} {len(made)} report(s)" + (f" (skipped mistakes: {made[0][0].skipped})" if made and made[0][0].skipped else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
