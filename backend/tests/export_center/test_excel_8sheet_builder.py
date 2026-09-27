import io
import openpyxl
import pandas as pd
import pytest

from app.export_center.builders.excel_builder import ExcelExportBuilder
from app.export_center.builders.pdf_builder import PdfExportBuilder
from app.export_center.models.export_models import BrandingConfig, ExportFormat, PageBudgetConfig, SmartMetricFilter
from app.pipeline.intelligence_pipeline import PowerPilotIntelligencePipeline


def test_excel_8sheet_audit_package_builder() -> None:
    df = pd.DataFrame(
        {
            "order_id": [1, 2, 3, 4],
            "customer_id": ["C101", "C102", "C103", "C104"],
            "sales": ["$150.00", "$250.00", "$350.00", "$450.00"],
            "profit": [30.0, 50.0, 70.0, 90.0],
            "order_date": ["2024-01-01", "2024-01-02", "2024-01-03", "2024-01-04"],
        }
    )
    pipeline = PowerPilotIntelligencePipeline()
    result = pipeline.run_pipeline(df, dataset_name="Retail_Sales_2024.csv")

    xlsx_bytes = ExcelExportBuilder.build_cleaned_excel(df, result)
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 0

    # Inspect Workbook Sheet Names & Structure
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    expected_sheets = [
        "Read Me",
        "Cleaned Dataset",
        "Cleaning Summary",
        "Data Quality Scorecard",
        "Original vs Cleaned",
        "Data Dictionary",
        "Audit Metadata",
        "Business Rule Validation",
    ]

    for sheet_name in expected_sheets:
        assert sheet_name in wb.sheetnames, f"Missing required sheet: {sheet_name}"


def test_smart_metric_filtering() -> None:
    cols = ["order_id", "customer_id", "Sales", "Profit", "Quantity", "row_num"]
    business_cols = SmartMetricFilter.filter_business_columns(cols)
    assert business_cols == ["Sales", "Profit", "Quantity"]


def test_pdf_page_budget_calculator() -> None:
    df = pd.DataFrame({"order_id": [1, 2], "sales": [10.0, 20.0]})
    pipeline = PowerPilotIntelligencePipeline()
    result = pipeline.run_pipeline(df, dataset_name="SmallTest.csv")

    budget = PdfExportBuilder.calculate_page_budget(result)
    assert budget.dataset_scale == "SMALL"
    assert budget.min_pages == 6
    assert budget.max_pages == 8


def test_cleaned_dataset_sheet_uses_a_native_table_not_per_cell_borders() -> None:
    """Regression guard for a real performance bug.

    The Cleaned Dataset sheet used to style every cell individually (a fresh
    Border object per cell). On a 9,800-row real-world dataset that loop alone
    took ~23 of ~27 total seconds to build the file, with no progress feedback
    in the UI -- exactly what makes an export feel broken rather than slow.

    A native Excel Table applies banded-row styling as one range-level
    definition instead of iterating every cell, and looks better besides
    (proper banding, built-in filter/sort). This asserts the table exists so a
    future edit can't silently reintroduce the per-cell loop.
    """
    df = pd.DataFrame(
        {
            "order_id": range(1, 5001),
            "region": (["North", "South", "East", "West"] * 1250),
            "sales": [round(100 + i * 1.37, 2) for i in range(5000)],
        }
    )
    pipeline = PowerPilotIntelligencePipeline()
    result = pipeline.run_pipeline(df, dataset_name="Large_Dataset.csv")

    import time

    started = time.perf_counter()
    xlsx_bytes = ExcelExportBuilder.build_cleaned_excel(df, result)
    elapsed = time.perf_counter() - started

    # Generous ceiling: the per-cell-border version took ~12s at this row count
    # (scales with rows x cols); the table-based version takes well under 1s.
    # 5s leaves headroom for slow CI runners while still catching a regression
    # to the old O(rows x cols) styling loop.
    assert elapsed < 5.0, (
        f"build_cleaned_excel took {elapsed:.2f}s for 5,000 rows -- "
        "this likely means per-cell styling was reintroduced on the data sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Cleaned Dataset"]
    assert "CleanedDataset" in ws.tables, (
        "expected a native Excel Table on the Cleaned Dataset sheet"
    )

    # Data integrity: the table-based path must still produce a correct sheet.
    assert ws.max_row == len(df) + 1
    assert ws.max_column == len(df.columns)
    header_values = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_values == list(df.columns)
