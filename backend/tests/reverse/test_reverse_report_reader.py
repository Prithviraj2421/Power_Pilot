"""Reading legacy reports: layouts, labels, number formats, derived cells and hostile files."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.reverse.errors import ReportReadError
from app.reverse.formulas import aggregate_of, evaluate, expand_range, references
from app.reverse.report_reader import read_report


def workbook(sheets: dict[str, list[list]], *, merges=(), formats=None, hidden=()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value is not None:
                    ws.cell(r, c, value)
        for ref, fmt in (formats or {}).items():
            ws[ref].number_format = fmt
        for merge in merges:
            ws.merge_cells(merge)
        if name in hidden:
            ws.sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def by_ref(report):
    return {t.cell_ref: t for t in report.targets}


# --- layouts ------------------------------------------------------------------------------


def test_region_by_year_table_with_a_title() -> None:
    rows = [
        ["Sales by Region and Year"],
        ["Region", 2023, 2024],
        ["East", 100.5, 200.25],
        ["West", 300.75, 400.0],
    ]
    report = read_report(workbook({"Summary": rows}), "r.xlsx")
    cells = by_ref(report)
    assert set(cells) == {"B3", "C3", "B4", "C4"}  # the year headers are labels, not numbers to explain
    east_2024 = cells["C3"]
    assert east_2024.value == 200.25 and east_2024.row_labels == ("East",) and east_2024.col_labels == ("2024",)
    assert east_2024.context_labels == ("Sales by Region and Year", "Summary")
    assert east_2024.row_axis_names == ("Region",)
    assert east_2024.id == "Summary!C3" and east_2024.full_precision


def test_generic_sheet_names_are_not_context() -> None:
    report = read_report(workbook({"Sheet1": [["A", "B"], ["x", 1.5]]}), "r.xlsx")
    assert report.targets[0].context_labels == ()


def test_merged_super_headers_and_two_header_rows() -> None:
    rows = [
        [None, 2023, None, 2024, None],
        ["Region", "Sales", "Profit", "Sales", "Profit"],
        ["East", 10.5, 1.5, 20.5, 2.5],
    ]
    report = read_report(workbook({"S": rows}, merges=["B1:C1", "D1:E1"]), "r.xlsx")
    cells = by_ref(report)
    assert cells["B3"].col_labels == ("2023", "Sales")
    assert cells["C3"].col_labels == ("2023", "Profit")
    assert cells["D3"].col_labels == ("2024", "Sales")
    assert cells["E3"].col_labels == ("2024", "Profit")


def test_super_headers_spread_over_blank_cells_when_there_are_no_merges() -> None:
    csv_text = ",2023,,2024,\nRegion,Sales,Profit,Sales,Profit\nEast,10.5,1.5,20.5,2.5\n"
    report = read_report(csv_text.encode(), "r.csv")
    assert [t.col_labels for t in report.targets] == [("2023", "Sales"), ("2023", "Profit"), ("2024", "Sales"), ("2024", "Profit")]


def test_pivot_style_groups_repeat_the_outer_label() -> None:
    rows = [
        ["Category", "Segment", "Sales"],
        ["Furniture", "Consumer", 10.5],
        [None, "Corporate", 20.5],
        [None, "Home Office", 30.5],
        ["Technology", "Consumer", 40.5],
        [None, "Corporate", 50.5],
    ]
    report = read_report(workbook({"S": rows}), "r.xlsx")
    assert [t.row_labels for t in report.targets] == [
        ("Furniture", "Consumer"),
        ("Furniture", "Corporate"),
        ("Furniture", "Home Office"),
        ("Technology", "Consumer"),
        ("Technology", "Corporate"),
    ]
    assert report.targets[0].row_axis_names == ("Category", "Segment")


def test_years_down_the_side_are_row_labels() -> None:
    rows = [["Year", "Sales"], [2023, 100.5], [2024, 200.5], [2025, 300.5]]
    report = read_report(workbook({"S": rows}), "r.xlsx")
    assert [(t.row_labels, t.value) for t in report.targets] == [(("2023",), 100.5), (("2024",), 200.5), (("2025",), 300.5)]


def test_a_numbering_column_is_neither_label_nor_data() -> None:
    rows = [["#", "Region", "Sales"], [1, "East", 10.5], [2, "West", 20.5], [3, "North", 30.5]]
    report = read_report(workbook({"S": rows}), "r.xlsx")
    assert [t.cell_ref for t in report.targets] == ["C2", "C3", "C4"]
    assert report.targets[0].row_labels == ("East",)


def test_two_tables_on_one_sheet_and_their_own_titles() -> None:
    rows = [
        ["Sales by Region"],
        ["Region", "Sales"],
        ["East", 10.5],
        ["West", 20.5],
        [],
        ["Orders by Segment"],
        ["Segment", "Orders"],
        ["Consumer", 7],
        ["Corporate", 9],
    ]
    report = read_report(workbook({"S": rows}), "r.xlsx")
    cells = by_ref(report)
    assert cells["B3"].context_labels[0] == "Sales by Region" and cells["B3"].row_labels == ("East",)
    assert cells["B8"].context_labels[0] == "Orders by Segment" and cells["B8"].col_labels == ("Orders",)


def test_a_title_separated_by_a_blank_row_still_applies() -> None:
    rows = [["Profit by Category"], [], ["Category", "Profit"], ["Furniture", 5.5], ["Technology", 6.5]]
    report = read_report(workbook({"S": rows}), "r.xlsx")
    assert report.targets[0].context_labels[0] == "Profit by Category"


def test_hidden_sheets_are_skipped_with_a_warning() -> None:
    sheets = {"Shown": [["A", "B"], ["x", 1.5]], "Secret": [["A", "B"], ["y", 2.5]]}
    report = read_report(workbook(sheets, hidden=("Secret",)), "r.xlsx")
    assert {t.sheet for t in report.targets} == {"Shown"} and any("Secret" in w for w in report.warnings)


# --- numbers as shown ---------------------------------------------------------------------


def test_spreadsheet_numbers_keep_the_precision_their_format_shows() -> None:
    rows = [["Metric", "Value"], ["Margin", 0.12345], ["Revenue", 1234567.891], ["Big", 1_500_000], ["Cost", -5.5], ["Raw", 32368.409999999996]]
    formats = {"B2": "0.0%", "B3": "#,##0", "B4": "#,##0.0,,", "B5": '"$"#,##0.00', "B6": "General"}
    cells = by_ref(read_report(workbook({"S": rows}, formats=formats), "r.xlsx"))

    assert (cells["B2"].value, cells["B2"].decimals_shown, cells["B2"].unit, cells["B2"].shown_text) == (0.12345, 1, "percent", "12.3%")
    assert cells["B2"].tolerance == pytest.approx(0.0005)
    assert (cells["B3"].decimals_shown, cells["B3"].tolerance) == (0, 0.5) and cells["B3"].shown_text == "1,234,568"
    assert cells["B4"].scale == 1e6 and cells["B4"].tolerance == pytest.approx(50_000) and cells["B4"].shown_text == "1.5M"
    assert (cells["B5"].unit, cells["B5"].decimals_shown) == ("currency", 2)
    assert cells["B6"].decimals_shown == 2  # floating-point noise is not "precision"


def test_csv_text_is_parsed_the_way_a_person_reads_it() -> None:
    csv_text = (
        "Item,Amount\n"
        'Revenue,"₹1,24,500"\n'
        'Loss,"(1,200)"\n'
        "Margin,12.5%\n"
        "Big,1.2M\n"
        "Note,see below\n"
    )
    cells = {t.row_labels[0]: t for t in read_report(csv_text.encode("utf-8"), "r.csv").targets}
    assert set(cells) == {"Revenue", "Loss", "Margin", "Big"}  # "see below" is text, not a number
    assert (cells["Revenue"].value, cells["Revenue"].unit) == (124500, "currency")
    assert cells["Loss"].value == -1200
    assert (cells["Margin"].value, cells["Margin"].unit, cells["Margin"].decimals_shown) == (0.125, "percent", 1)
    assert cells["Big"].value == 1_200_000 and cells["Big"].tolerance == pytest.approx(50_000)
    assert not any(t.full_precision for t in cells.values())  # CSV only has the shown text


def test_csv_dialects_and_encodings() -> None:
    semicolons = "Region;Sales\nEast;1.234,50\n".encode("cp1252")
    assert read_report(semicolons, "r.csv").targets[0].value == 1234.5
    with_bom = "﻿Region,Sales\nEast,10.5\n".encode("utf-8")
    assert read_report(with_bom, "r.csv").targets[0].row_labels == ("East",)


# --- derived cells ------------------------------------------------------------------------


def test_formula_totals_are_derived_even_without_a_saved_result() -> None:
    rows = [
        ["Region", "Sales"],
        ["East", 10.5],
        ["West", 20.5],
        ["Total", "=SUM(B2:B3)"],
        ["Twice", "=B4*2"],
    ]
    cells = by_ref(read_report(workbook({"S": rows}), "r.xlsx"))  # openpyxl saves no cached results
    total = cells["B4"]
    assert total.value == 31.0 and total.derived.kind == "formula"
    assert (total.derived.op, total.derived.sources) == ("sum", ("S!B2", "S!B3"))
    assert cells["B5"].value == 62.0 and cells["B5"].derived.op is None  # derived, but not a plain aggregate
    assert cells["B2"].derived is None


def test_a_formula_reading_another_sheet_is_an_ordinary_number() -> None:
    sheets = {"Data": [["x", 5]], "S": [["Item", "Value"], ["Copy", "=Data!B1*2"], ["Real", 7.5]]}
    report = read_report(workbook(sheets), "r.xlsx")
    cells = by_ref(report)
    assert "B2" not in cells and any("S!B2" in w and "Data!B1" in w for w in report.warnings)  # no saved result to explain
    assert cells["B3"].value == 7.5


def test_typed_total_rows_and_columns_are_derived_from_what_they_sit_beside() -> None:
    csv_text = "Region,2023,2024,Total\nEast,10,20,30\nWest,1,2,3\nTotal,11,22,33\nAll Regions,11,22,33\n"
    cells = {(t.row_labels[0], t.col_labels[-1]): t for t in read_report(csv_text.encode(), "r.csv").targets}
    assert cells[("East", "2023")].derived is None
    assert cells[("East", "Total")].derived.sources == ("r!B2", "r!C2")  # the cells to its left in the same row
    column_total = cells[("Total", "2023")].derived
    assert column_total.kind == "total_label" and column_total.op == "sum" and len(column_total.sources) == 2
    assert cells[("Total", "Total")].derived is not None
    assert cells[("All Regions", "2023")].derived is None  # "All" is not a sum we can name: the data has to explain it


def test_subtotals_stop_at_the_previous_subtotal_and_grand_total_skips_them() -> None:
    csv_text = "Item,Amount\nA,1\nB,2\nSubtotal,3\nC,4\nD,5\nSubtotal,9\nGrand Total,12\n"
    cells = {t.row_labels[0] + str(t.cell_ref): t for t in read_report(csv_text.encode(), "r.csv").targets}
    first, second, grand = cells["SubtotalB4"], cells["SubtotalB7"], cells["Grand TotalB8"]
    assert len(first.derived.sources) == 2 and len(second.derived.sources) == 2 and len(grand.derived.sources) == 4


# --- the formula evaluator ----------------------------------------------------------------


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("=SUM(A1:A3)", 6.0),
        ("=AVERAGE(A1:A3)", 2.0),
        ("=MIN(A1:A3)+MAX(A1:A3)", 4.0),
        ("=A1+A2*A3", 7.0),
        ("=(A1+A2)*A3", 9.0),
        ("=-A1+10", 9.0),
        ("=A3/A1", 3.0),
        ("=2^3", 8.0),
        ("=50%*A3", 1.5),
        ("=ROUND(A3/7, 2)", 0.43),
        ("=ABS(0-A2)", 2.0),
        ("=COUNT(A1:A3)", 3.0),
        ("=SUM(A1,A2,10)", 13.0),
        ("=$A$1+$A$2", 3.0),
        ("=sum(a1:a2)", 3.0),
    ],
)
def test_evaluator(formula, expected) -> None:
    cells = {"A1": 1.0, "A2": 2.0, "A3": 3.0}
    assert evaluate(formula, cells.get) == pytest.approx(expected)


@pytest.mark.parametrize("formula", ["=Other!A1+1", '=IF(A1>1,1,2)', "=A1/0", "=VLOOKUP(A1,B:C,2)", '="x"&A1', "=SUM(", "=A1+"])
def test_formulas_that_are_not_understood_give_none_not_a_guess(formula) -> None:
    assert evaluate(formula, {"A1": 1.0}.get) is None
    assert references(formula) is None


def test_references_and_aggregates() -> None:
    assert references("=SUM(B2:B4)") == ["B2", "B3", "B4"]
    assert references("=A1+A1+B2") == ["A1", "B2"]
    assert expand_range("A1:B2") == ["A1", "B1", "A2", "B2"]
    assert aggregate_of("=SUM(B2:B4)") == "sum" and aggregate_of("=AVERAGE(B2:B4)") == "average"
    assert aggregate_of("=B2+B3+B4") == "sum" and aggregate_of("=B2*2") is None and aggregate_of("=SUM(B2:B4)/2") is None


# --- pdf ----------------------------------------------------------------------------------


def pdf_table(rows: list[list[str]], title: str = "Sales report") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    table = Table(rows)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(buffer, pagesize=A4).build([Paragraph(title, getSampleStyleSheet()["Title"]), table])
    return buffer.getvalue()


def test_a_pdf_table_is_read_like_any_other_grid() -> None:
    pdf = pdf_table([["Region", "2023", "2024"], ["East", "1,234.50", "(200)"], ["West", "12.5%", "Rs 1,24,500"]])
    report = read_report(pdf, "report.pdf")
    cells = {(t.row_labels[0], t.col_labels[0]): t for t in report.targets}
    assert cells[("East", "2023")].value == 1234.5 and cells[("East", "2024")].value == -200
    assert (cells[("West", "2023")].value, cells[("West", "2023")].unit) == (0.125, "percent")
    assert cells[("West", "2024")].value == 124500
    assert report.targets[0].sheet == "Page 1" and not report.targets[0].full_precision


# --- hostile and unsupported files --------------------------------------------------------


@pytest.mark.parametrize(
    ("content", "name", "message"),
    [
        (b"", "r.xlsx", "empty"),
        (b"hello", "r.txt", "Unsupported file type"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100, "old.xls", r"\.xls"),
        (b"not a zip file at all", "r.xlsx", "could not be opened"),
        (b"%PDF-1.4 garbage", "r.pdf", "could not be read"),
        (b"Region,Sales\nEast,abc\n", "r.csv", "No numbers"),
        (b"just,words\nhere,only\n", "r.csv", "No numbers"),
    ],
)
def test_unreadable_files_get_a_plain_message(content, name, message) -> None:
    with pytest.raises(ReportReadError, match=message):
        read_report(content, name)


def test_a_pdf_without_tables_says_so() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4).build([Paragraph("Just a paragraph of words.", getSampleStyleSheet()["Normal"])])
    with pytest.raises(ReportReadError, match="No tables"):
        read_report(buffer.getvalue(), "r.pdf")


def test_a_data_table_is_refused_as_not_a_report() -> None:
    lines = ["Name,Amount"] + [f"row{i},{i}.5" for i in range(3_100)]
    with pytest.raises(ReportReadError, match="data table"):
        read_report("\n".join(lines).encode(), "big.csv")


def test_a_workbook_with_only_hidden_sheets_is_refused() -> None:
    import zipfile

    source = workbook({"S": [["A", "B"], ["x", 1.5]], "T": [["A", "B"], ["y", 2.5]]})
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as zin, zipfile.ZipFile(out, "w") as zout:  # openpyxl will not save this
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = data.replace(b"<sheet ", b'<sheet state="hidden" ')
            zout.writestr(item, data)
    with pytest.raises(ReportReadError, match="could not be opened"):
        read_report(out.getvalue(), "r.xlsx")
