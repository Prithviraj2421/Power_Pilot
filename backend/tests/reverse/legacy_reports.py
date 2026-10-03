"""Builds a legacy report from the golden Superstore data, with known formulas and planted mistakes.

The report is generated at test time (no binary fixtures) so the expected formula of every number is
known exactly, and the same content can be written as .xlsx, .csv and .pdf.

Tables (all from ``tests/golden/superstore.csv``):

* Sales by Region and Year   SUM(Sales) where Region = r and YEAR(Order Date) = y, an "All Regions" row
                             typed from the data, and a Total row (a formula in the .xlsx)
* Profit margin % by Category   SUM(Profit) / SUM(Sales), 1 decimal
* Customers by Segment       DISTINCT COUNT(Customer ID)   (Customer Name fits some of them too)
* Average discount by Ship Mode   AVERAGE(Discount), 1 decimal

Three mistakes are planted in the first table: a swapped-digit typo (East 2024), one order left out of
a number (Central 2023), and a typed "All Regions" total for 2024 that forgot a whole region (West).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import pandas as pd

from app.intelligence.kpi.compilers import PandasCompiler
from app.intelligence.kpi.ir import Compare, DatePart, Expr, Filter, Measure, Op, Ratio
from app.reverse.describe import canonical_filters

GOLDEN = Path(__file__).resolve().parents[1] / "golden" / "superstore.csv"
REGIONS = ["Central", "East", "South", "West"]
YEARS = [2023, 2024]
CATEGORIES = ["Furniture", "Office Supplies", "Technology"]
SEGMENTS = ["Consumer", "Corporate", "Home Office"]
SHIP_MODES = ["First Class", "Same Day", "Second Class", "Standard Class"]

PC = PandasCompiler()


def golden() -> pd.DataFrame:
    return pd.read_csv(GOLDEN)


def year_filter(year: int) -> Filter:
    return Filter("Order Date", Compare.EQ, year, DatePart.YEAR)


def sales_by(region: Optional[str], year: int) -> Expr:
    filters = ([Filter("Region", Compare.EQ, region)] if region else []) + [year_filter(year)]
    return Measure(Op.SUM, "Sales", canonical_filters(tuple(filters)))


@dataclass
class Cell:
    """One number in the report and the formula that really produced it (None for planted mistakes)."""

    table: str
    row: str
    col: str
    value: float
    expected: Optional[Expr]  # None when the number is a planted mistake
    fmt: str = "#,##0.00"
    mistake: Optional[str] = None  # "swapped-digits" | "one-order" | "missing-region"
    formula: Optional[str] = None  # spreadsheet formula (xlsx only) for derived cells


@dataclass
class Report:
    df: pd.DataFrame
    tables: dict[str, list[list]] = field(default_factory=dict)  # table title -> grid rows (text and Cell objects)
    cells: list[Cell] = field(default_factory=list)

    def cell(self, table: str, row: str, col: str) -> Cell:
        return next(c for c in self.cells if (c.table, c.row, c.col) == (table, row, col))


T1 = "Sales by Region and Year"
T2 = "Profit margin % by Category"
T3 = "Customers by Segment"
T4 = "Average discount by Ship Mode"


def build(df: Optional[pd.DataFrame] = None, *, mistakes: bool = True) -> Report:
    df = golden() if df is None else df
    report = Report(df)

    def true(expr: Expr) -> float:
        return float(PC.evaluate(expr, df))

    # --- table 1 -----------------------------------------------------------------------
    rows: list[list] = [[T1], ["Region", *YEARS]]
    for region in REGIONS:
        line: list = [region]
        for year in YEARS:
            expr = sales_by(region, year)
            cell = Cell(T1, region, str(year), round(true(expr), 2), expr)
            if mistakes and (region, year) == ("East", 2024):
                cell.mistake, cell.expected = "swapped-digits", None
                digits = f"{cell.value:.2f}"
                cell.value = float(digits[:3] + digits[4] + digits[3] + digits[5:])  # 16079.19 -> 16097.19
            if mistakes and (region, year) == ("Central", 2023):
                in_scope = df[(df["Region"] == "Central") & (pd.to_datetime(df["Order Date"], dayfirst=True).dt.year == 2023)]
                orders = in_scope.groupby("Order ID")["Sales"].agg(["sum", "size"])
                order_id = orders[orders["size"] >= 2]["sum"].index[0]
                cell.mistake, cell.expected = "one-order", None
                cell.value = round(cell.value - float(orders.loc[order_id, "sum"]), 2)
                cell.order_id = order_id  # type: ignore[attr-defined]
            report.cells.append(cell)
            line.append(cell)
        rows.append(line)

    line = ["All Regions"]
    for year in YEARS:
        expr = sales_by(None, year)
        cell = Cell(T1, "All Regions", str(year), round(true(expr), 2), expr)
        if mistakes and year == 2024:
            west = true(sales_by("West", 2024))
            cell.mistake, cell.expected = "missing-region", None
            cell.value = round(cell.value - west, 2)
        report.cells.append(cell)
        line.append(cell)
    rows.append(line)

    total = ["Total"]
    for index, year in enumerate(YEARS):
        parts = [report.cell(T1, region, str(year)).value for region in REGIONS]
        column = "BC"[index]
        cell = Cell(T1, "Total", str(year), round(sum(parts), 2), None, formula=f"=SUM({column}3:{column}6)")
        report.cells.append(cell)
        total.append(cell)
    rows.append(total)
    report.tables[T1] = rows

    # --- table 2 -----------------------------------------------------------------------
    rows = [[T2], ["Category", "Profit margin %"]]
    for category in CATEGORIES:
        scope = (Filter("Category", Compare.EQ, category),)
        expr = Ratio(Measure(Op.SUM, "Profit", scope), Measure(Op.SUM, "Sales", scope))
        cell = Cell(T2, category, "Profit margin %", true(expr), expr, fmt="0.0%")
        report.cells.append(cell)
        rows.append([category, cell])
    report.tables[T2] = rows

    # --- table 3 -----------------------------------------------------------------------
    rows = [[T3], ["Segment", "Customers"]]
    for segment in SEGMENTS:
        expr = Measure(Op.DISTINCT_COUNT, "Customer ID", (Filter("Segment", Compare.EQ, segment),))
        cell = Cell(T3, segment, "Customers", true(expr), expr, fmt="0")
        report.cells.append(cell)
        rows.append([segment, cell])
    report.tables[T3] = rows

    # --- table 4 -----------------------------------------------------------------------
    rows = [[T4], ["Ship Mode", "Average discount"]]
    for mode in SHIP_MODES:
        expr = Measure(Op.AVERAGE, "Discount", (Filter("Ship Mode", Compare.EQ, mode),))
        cell = Cell(T4, mode, "Average discount", true(expr), expr, fmt="0.0%")
        report.cells.append(cell)
        rows.append([mode, cell])
    report.tables[T4] = rows
    return report


def shown(cell: Cell) -> str:
    """The text a person would read (CSV and PDF only carry this)."""
    if cell.fmt == "0.0%":
        return f"{cell.value * 100:.1f}%"
    if cell.fmt == "0":
        return f"{cell.value:.0f}"
    return f"{cell.value:,.2f}"


def _text_rows(report: Report) -> list[tuple[str, list[list[str]]]]:
    out = []
    for title, rows in report.tables.items():
        out.append((title, [[shown(x) if isinstance(x, Cell) else str(x) for x in row] for row in rows]))
    return out


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    r = 1
    for rows in report.tables.values():
        for row in rows:
            for c, item in enumerate(row, start=1):
                if isinstance(item, Cell):
                    ws.cell(r, c, item.formula or item.value).number_format = item.fmt
                else:
                    ws.cell(r, c, item)
            r += 1
        r += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_csv(report: Report) -> bytes:
    lines = []
    for _, rows in _text_rows(report):
        for row in rows:
            lines.append(",".join(f'"{x}"' if "," in x else x for x in row))
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def to_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

    flow = []
    for _, rows in _text_rows(report):
        width = max(len(r) for r in rows)
        padded = [r + [""] * (width - len(r)) for r in rows]
        table = Table(padded)
        style = [("GRID", (0, 0), (-1, -1), 0.5, colors.black)]
        if width > 1:
            style.append(("SPAN", (0, 0), (-1, 0)))  # the title row spans the table
        table.setStyle(TableStyle(style))
        flow += [table, Spacer(1, 18)]
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4).build(flow)
    return buffer.getvalue()
