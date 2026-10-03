"""Read a legacy report (.xlsx / .csv / .pdf) into the numbers that need explaining.

Spreadsheet numbers are read at the precision the cell *shows* (its number format) while the stored
full-precision value is kept; CSV and PDF only have the text, so that is all we can match against.
Formula cells that read this sheet are marked derived; so are "Total" rows and columns.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import PurePath
from typing import Optional

from app.reverse import formulas
from app.reverse.errors import ReportReadError
from app.reverse.grid import Cell, Grid, extract
from app.reverse.models import ParsedReport
from app.reverse.numbers import ParsedNumber, interpret_format, is_date_format, parse_number, show

MAX_CELLS = 400_000  # cells in one worksheet / file we are willing to scan
MAX_TARGETS = 3_000  # more numbers than this is a data table, not a report
MAX_PDF_PAGES = 30
SUPPORTED = (".xlsx", ".xlsm", ".csv", ".pdf")


def read_report(content: bytes, filename: str) -> ParsedReport:
    """Parse an uploaded report. Raises ``ReportReadError`` with a message fit to show the user."""
    suffix = PurePath(filename or "").suffix.lower()
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or suffix == ".xls":
        raise ReportReadError("Old .xls files are not supported. Open it in Excel and save it as .xlsx, then upload again.")
    if suffix not in SUPPORTED:
        raise ReportReadError(f"Unsupported file type '{suffix or filename}'. Upload an .xlsx, .csv or .pdf report.")
    if not content:
        raise ReportReadError("The file is empty.")

    if suffix in (".xlsx", ".xlsm"):
        grids, warnings = _read_xlsx(content)
    elif suffix == ".pdf":
        grids, warnings = _read_pdf(content)
    else:
        grids, warnings = [_read_csv(content, PurePath(filename).stem or "Report")], []

    report = ParsedReport(warnings=warnings)
    for grid in grids:
        targets, layout = extract(grid)
        report.targets.extend(targets)
        report.layout.append(layout)
    if not report.targets:
        raise ReportReadError(
            "No numbers were found in the report. PowerPilot needs a table with a header row and the numbers under it."
        )
    if len(report.targets) > MAX_TARGETS:
        raise ReportReadError(
            f"The file holds {len(report.targets):,} numbers, which looks like a data table rather than a finished report. "
            "Upload the report itself (the summary tables people read)."
        )
    return report


# --- xlsx ---------------------------------------------------------------------------------


def _date_label(value: object, number_format: str) -> str:
    moment = value if isinstance(value, dt.datetime) else dt.datetime.combine(value, dt.time())  # type: ignore[arg-type]
    fmt = (number_format or "").lower()
    if "d" in fmt:
        return moment.strftime("%Y-%m-%d")
    if "m" in fmt and "y" in fmt:
        return moment.strftime("%b %Y")
    return moment.strftime("%Y") if "y" in fmt else moment.strftime("%Y-%m-%d")


def _read_xlsx(content: bytes) -> tuple[list[Grid], list[str]]:
    from openpyxl import load_workbook

    try:
        with_values = load_workbook(io.BytesIO(content), data_only=True)
        with_formulas = load_workbook(io.BytesIO(content), data_only=False)
    except Exception as exc:  # openpyxl raises many things for bad, encrypted or truncated files
        raise ReportReadError(
            "The Excel file could not be opened. It may be password-protected, damaged, or not really an .xlsx file."
        ) from exc

    grids: list[Grid] = []
    warnings: list[str] = []
    for ws in with_formulas.worksheets:
        if ws.sheet_state != "visible":
            warnings.append(f"Sheet '{ws.title}' is hidden and was skipped.")
            continue
        if (ws.max_row or 0) * (ws.max_column or 0) > MAX_CELLS:
            raise ReportReadError(f"Sheet '{ws.title}' is too large to read as a report ({ws.max_row:,} rows).")
        cached = with_values[ws.title]
        grid = Grid(ws.title)
        grid.merges = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]

        numbers: dict[str, float] = {}  # plain numeric cells, for evaluating formulas that have no saved result
        pending: list[tuple[int, int, str, str, str]] = []
        for row in ws.iter_rows():
            for source in row:
                value = source.value
                if value is None:
                    continue
                r, c = source.row - 1, source.column - 1
                fmt = source.number_format
                if source.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    result = cached[source.coordinate].value
                    if isinstance(result, (int, float)) and not isinstance(result, bool):
                        grid.add(_numeric(r, c, float(result), fmt, formula=str(value)))
                        numbers[source.coordinate] = float(result)
                    else:
                        pending.append((r, c, source.coordinate, str(value), fmt))  # type: ignore[arg-type]
                    continue
                if isinstance(value, bool):
                    grid.add(Cell(r, c, str(value).upper()))
                elif isinstance(value, (int, float)):
                    if is_date_format(fmt):
                        continue  # a serial number displayed as a date, but read as a plain number: skip
                    grid.add(_numeric(r, c, float(value), fmt))
                    numbers[source.coordinate] = float(value)
                elif isinstance(value, (dt.datetime, dt.date)):
                    grid.add(Cell(r, c, _date_label(value, fmt)))
                else:
                    text = str(value).strip()
                    number = parse_number(text)
                    grid.add(Cell(r, c, text, number))
                    if number is not None:
                        numbers[source.coordinate] = number.value

        for _ in range(10):  # formulas that read other formulas settle in a few passes
            still: list[tuple[int, int, str, str, str]] = []
            for r, c, coordinate, formula, fmt in pending:
                result = formulas.evaluate(formula, numbers.get)
                if result is None:
                    still.append((r, c, coordinate, formula, fmt))
                    continue
                grid.add(_numeric(r, c, result, fmt, formula=formula))  # type: ignore[arg-type]
                numbers[coordinate] = result
            if len(still) == len(pending):
                break
            pending = still
        for r, c, coordinate, formula, _fmt in pending:
            warnings.append(
                f"{ws.title}!{coordinate} is a formula ({formula}) with no saved result that PowerPilot cannot work out, so it was skipped."
            )
        grids.append(grid)
    if not grids:
        raise ReportReadError("The workbook has no visible sheets.")
    return grids, warnings


def _numeric(row: int, col: int, value: float, number_format: str, formula: Optional[str] = None) -> Cell:
    info = interpret_format(number_format, value)
    return Cell(
        row,
        col,
        text=show(value, info),
        number=ParsedNumber(value, info.decimals, info.unit, info.scale),
        formula=formula,
        full_precision=True,
    )


# --- csv ----------------------------------------------------------------------------------


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _read_csv(content: bytes, name: str) -> Grid:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if sum(len(r) for r in rows) > MAX_CELLS:
        raise ReportReadError("The CSV is too large to read as a report.")
    return _text_grid(name, rows)


def _text_grid(name: str, rows: list[list[Optional[str]]]) -> Grid:
    grid = Grid(name)
    for r, row in enumerate(rows):
        for c, raw in enumerate(row):
            text = re.sub(r"\s+", " ", raw or "").strip()
            if text:
                grid.add(Cell(r, c, text, parse_number(text)))
    return grid


# --- pdf ----------------------------------------------------------------------------------


def _read_pdf(content: bytes) -> tuple[list[Grid], list[str]]:
    import pdfplumber

    grids: list[Grid] = []
    warnings: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = pdf.pages
            if len(pages) > MAX_PDF_PAGES:
                warnings.append(f"Only the first {MAX_PDF_PAGES} of {len(pages)} pages were read.")
            for number, page in enumerate(pages[:MAX_PDF_PAGES], start=1):
                tables = page.extract_tables() or page.extract_tables(
                    {"vertical_strategy": "text", "horizontal_strategy": "text"}
                )
                for index, table in enumerate(t for t in tables if t and len(t) >= 2):
                    label = f"Page {number}" if len(tables) == 1 else f"Page {number} table {index + 1}"
                    grids.append(_text_grid(label, table))
    except ReportReadError:
        raise
    except Exception as exc:  # pdfminer raises a family of errors for encrypted or damaged files
        raise ReportReadError(
            "The PDF could not be read. It may be password-protected, damaged, or a scan (PowerPilot cannot read images of tables)."
        ) from exc
    if not grids:
        raise ReportReadError(
            "No tables were found in the PDF. PowerPilot reads tables with real text; a scanned image needs to be typed or exported to Excel first."
        )
    return grids, warnings
