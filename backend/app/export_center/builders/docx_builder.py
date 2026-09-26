import io
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.export_center.models.export_models import BrandingConfig, SmartMetricFilter
from app.models.master_intelligence_result import MasterIntelligenceResult


class DocxExportBuilder:
    """
    Builder generating editable executive Word reports (.docx) using python-docx.
    """

    @staticmethod
    def build_executive_docx(result: MasterIntelligenceResult, branding: BrandingConfig = BrandingConfig()) -> bytes:
        doc = Document()

        # Custom Styles
        title_style = doc.styles["Title"]
        title_style.font.name = "Calibri"
        title_style.font.size = Pt(26)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        h1_style = doc.styles["Heading 1"]
        h1_style.font.name = "Calibri"
        h1_style.font.size = Pt(18)
        h1_style.font.bold = True
        h1_style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        # Title Page / Cover Section
        title_p = doc.add_paragraph("Executive Intelligence & BI Report", style="Title")
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph(f"Prepared for: {branding.prepared_for} | By: {branding.prepared_by} | {branding.report_version}")
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # Section 1: Executive Summary
        doc.add_heading("1. Executive Summary & Strategic Synthesis", level=1)
        dataset_name = result.dataset_profile.dataset_name
        domain = result.dataset_profile.detected_domain.upper()
        doc.add_paragraph(
            f"This executive intelligence report synthesizes structural, semantic, statistical, and strategic findings for dataset '{dataset_name}' categorized under the {domain} business domain."
        )

        summary_obj = result.insight_report.executive_summary if result.insight_report else None
        if summary_obj:
            overview_text = getattr(summary_obj, "overview", getattr(summary_obj, "headline", "Executive Analysis Overview"))
            doc.add_paragraph(f"Overview: {overview_text}")

        # Section 2: Data Quality & Preparation Audit
        doc.add_heading("2. Stage 1 Data Quality Audit", level=1)
        if result.quality_report:
            doc.add_paragraph(
                f"Overall Dataset Quality Score: {result.quality_report.overall_score:.1f}% (Quality Grade: {result.quality_report.grade.value})"
            )
            doc.add_paragraph(result.quality_report.grade_explanation)

        # Section 3: Recommended DAX Measures
        doc.add_heading("3. Executable DAX Measures Studio", level=1)
        if result.kpi_report:
            all_kpis = list(result.kpi_report.primary_kpis) + list(result.kpi_report.secondary_kpis)
            for kpi in all_kpis[:8]:
                doc.add_paragraph(f"KPI Measure: {kpi.name} ({kpi.priority.value if hasattr(kpi.priority, 'value') else str(kpi.priority)})")
                code_p = doc.add_paragraph(kpi.formula)
                code_p.paragraph_format.left_indent = Inches(0.5)

        # Section 4: Strategic Decisions & Next Steps
        doc.add_heading("4. Strategic Management Actions", level=1)
        if result.decision_report:
            for dec in result.decision_report.primary_decisions[:5]:
                title = getattr(dec, "action_title", getattr(dec, "title", "Decision Action"))
                doc.add_paragraph(f"• {title}: {dec.description}", style="List Bullet")

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
